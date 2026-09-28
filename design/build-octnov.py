"""Build October and November from strategy/plan-*.json.

    python3 design/build-octnov.py        # then: python3 design/render.py octnov

The plan files are the only thing to hand edit. From them this writes:

  design/octnov-data.js         slide copy for carousel.html and single.html
  content/<date>-<slug>/        one folder per post, with brief.md
  strategy/gbp-daily-content.json   Oct 1 onward replaced, Mon to Thu only
  strategy/gbp-photos.json      any newly fetched GBP photos appended
  strategy/ig-content-strategy.xlsx
      Posting Schedule   Oct 1 onward replaced, two SEO columns added
      Reels & Video      Friday and Sunday slots moved onto workdays
      GBP Daily          rebuilt by build-gbp-daily.py
      How to Use, Monthly Themes, Hashtag Sets, Ideas Backlog   refreshed

Nothing is ever scheduled Friday, Saturday or Sunday from Oct 1: those are
days off. The script refuses a plan that tries.

Photos come from design/fetch-octnov-photos.py. Until a post's photo is on
disk its artwork renders in the type-only version of the same layout and its
row reads Design rather than Ready. Re-run this, then render, once they land.

Safe to re-run. Status and the tracking columns typed into Posting Schedule
are kept, matched by date.
"""
import datetime
import glob
import json
import os
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESIGN = os.path.join(ROOT, 'design')
SHEET = os.path.join(ROOT, 'strategy', 'ig-content-strategy.xlsx')
PICKS = os.path.join(ROOT, 'strategy', 'octnov-photos.json')
GBP_JSON = os.path.join(ROOT, 'strategy', 'gbp-daily-content.json')
GBP_PHOTOS = os.path.join(ROOT, 'strategy', 'gbp-photos.json')

REPO = 'RepoKing23/InstagramCarousel'
BRANCH = 'claude/serene-curie-xunco2'
START = '2026-10-01'
WEEK_ONE = datetime.date(2026, 8, 17)          # Monday of week 1 in the planner
TREE = f'https://github.com/{REPO}/tree/{BRANCH}'
BLOB = f'https://github.com/{REPO}/blob/{BRANCH}'
RAW = f'https://raw.githubusercontent.com/{REPO}/{BRANCH}'

SET_F = ('#electrolysisoakville #oakvilleelectrolysis #facialelectrolysis '
         '#electrolysis #permanenthairremoval #hairremovaloakville '
         '#burlingtonelectrolysis #miltonontario #mississaugabeauty '
         '#haltonregion #luxurybeautybycleor #npcleo')

SEO_COLS = ['SEO Keyword (GSC)', 'Alt Text']

# Search Console, Performance on Search, last 28 days to 2026-09-28, web,
# luxurybeautyaesthetic.com. Every post from Oct 1 targets one of these.
GSC = {
    'botox therapy treatment oakville': (0, 13, 54.85),
    'cleo oakville': (0, 3, 1.67),
    'electrolysis oakville': (0, 2, 4.5),
    'facial electrolysis': (0, 1, 21.0),
    'luxury spa oakville': (0, 1, 25.0),
    'dermal fillers': (0, 1, 41.0),
    'oakville botox': (0, 1, 57.0),
}


def gsc_label(kw):
    clicks, impr, pos = GSC[kw]
    return f'{kw} (GSC: {impr} impr, pos {pos:g})'


# --- plan -----------------------------------------------------------------

def load_plan():
    posts = []
    for f in sorted(glob.glob(os.path.join(ROOT, 'strategy', 'plan-*.json'))):
        posts += json.load(open(f, encoding='utf-8'))['posts']
    posts.sort(key=lambda p: p['date'])
    for p in posts:
        d = datetime.date.fromisoformat(p['date'])
        if d.weekday() > 3:
            sys.exit(f"{p['date']} is a {d:%A}. Friday to Sunday are off, move it.")
        p['_d'] = d
        p['folder'] = f"{p['date']}-{p['slug']}"
    return posts


def load_picks():
    if os.path.exists(PICKS):
        return json.load(open(PICKS, encoding='utf-8'))['picks']
    return {}


def ig_photo(p, picks):
    """Relative path the templates load, and the Unsplash pick behind it."""
    rel = f"photos/octnov/{p['date']}-ig.jpg"
    if os.path.exists(os.path.join(DESIGN, rel)):
        return rel, picks.get(f"{p['date']}-ig")
    return None, None


def gbp_pick(p, picks):
    pk = picks.get(f"{p['date']}-gbp")
    if pk and os.path.exists(os.path.join(DESIGN, 'photos', 'gbp', f"{pk['id']}.jpg")):
        return pk
    return None


def slide_files(p):
    """The finished images for a post, in posting order."""
    if p['kind'] == 'result':
        return list(p['result']['files'])
    if p['kind'] == 'carousel':
        return [f'slide-{i}.jpg' for i in range(1, len(p['ig']['points']) + 3)]
    return ['post.jpg']


def git(*args):
    subprocess.run(['git', '-C', ROOT, *args], check=True)


# --- content folders -------------------------------------------------------

def move_results(posts):
    """Real client results already designed keep their artwork; only their
    date moved, because Friday and Saturday are no longer posting days."""
    for p in posts:
        r = p.get('result')
        if not r:
            continue
        old = os.path.join(ROOT, 'content', r['folder'])
        new = os.path.join(ROOT, 'content', p['folder'])
        if old != new and os.path.isdir(old) and not os.path.exists(new):
            git('mv', old, new)
            print(f"  moved {r['folder']} -> {p['folder']}")
        cover = p.get('cover_file')
        if cover:
            src = os.path.join(new, 'images', cover)
            if os.path.exists(src):
                git('mv', src, os.path.join(new, 'images', 'slide-1.jpg'))
        r['files'] = ['slide-1.jpg' if f == cover else f for f in r['files']]
        if p['kind'] == 'carousel':
            r['files'] = slide_files({'kind': 'carousel', 'ig': p['ig']})


def prune_old(posts):
    """Oct and Nov folders from the old plan that nothing points at now. Only
    ever removes a folder holding nothing but its brief."""
    keep = {p['folder'] for p in posts}
    for path in sorted(glob.glob(os.path.join(ROOT, 'content', '2026-1[0-2]-*'))):
        name = os.path.basename(path)
        if name in keep:
            continue
        files = [os.path.relpath(os.path.join(dp, f), path)
                 for dp, _, fs in os.walk(path) for f in fs]
        if set(files) <= {'brief.md', os.path.join('images', '.gitkeep')}:
            git('rm', '-rq', path)
            shutil.rmtree(path, ignore_errors=True)
            print(f'  removed {name} (old plan, no artwork)')
        else:
            print(f'  KEPT {name}: has artwork but is not in the plan')


def prune_gbp(posts):
    """Old daily cards for Oct onward, replaced by the Mon to Thu set."""
    keep = {os.path.basename(gbp_image(p)) for p in posts}
    for path in sorted(glob.glob(os.path.join(ROOT, 'content', 'gbp', '2026-1[0-2]-*.jpg'))):
        if os.path.basename(path) not in keep:
            git('rm', '-q', path)
            print(f'  removed content/gbp/{os.path.basename(path)} (old daily card)')


def gbp_image(p):
    return f"content/gbp/{p['date']}-{p['slug']}.jpg"


# --- template data --------------------------------------------------------

def small(rel):
    """Photos only held at Unsplash small size go in an arch frame, not full bleed."""
    from PIL import Image
    with Image.open(os.path.join(DESIGN, rel)) as im:
        return im.width < 1000


def write_data_js(posts, picks):
    carousels, singles = {}, {}
    for p in posts:
        ig = p.get('ig')
        if not ig:
            continue
        photo, _ = ig_photo(p, picks)
        if p['kind'] == 'carousel':
            entry = {k: ig[k] for k in ('label', 'coverSerif', 'coverScript', 'coverSub',
                                        'points', 'ctaSerif', 'ctaScript', 'ctaBody', 'button')}
            if photo:
                key = 'coverFrame' if small(photo) else 'coverPhoto'
                entry.update({key: photo, 'coverFocal': 'center 30%'})
            carousels[p['slug']] = entry
        else:
            entry = {'label': ig['label'], 'serif': ig['serif'], 'script': ig['script'],
                     'body': ig['body'], 'cta': ig['cta'], 'theme': ig.get('theme', 'light')}
            # photo led when there is a photo, the plain statement card until then
            entry['kind'] = ('frame' if small(photo) else 'photo') if photo else 'statement'
            if photo:
                entry.update(photo=photo, focal='center 30%')
            singles[p['slug']] = entry
    out = os.path.join(DESIGN, 'octnov-data.js')
    with open(out, 'w', encoding='utf-8') as f:
        f.write('/* Generated by design/build-octnov.py from strategy/plan-*.json.\n'
                '   Edit the plan, not this file. */\n')
        f.write('window.OCTNOV = ')
        json.dump({'carousels': carousels, 'singles': singles}, f, indent=1, ensure_ascii=False)
        f.write(';\n')
    print(f'wrote design/octnov-data.js  ({len(carousels)} carousels, {len(singles)} singles)')


# --- briefs ---------------------------------------------------------------

FORMAT = {'carousel': 'Carousel', 'single': 'Single Image', 'result': 'Before/After'}
PLATFORMS = {'carousel': 'Instagram feed + GBP', 'single': 'IG + FB + TikTok photo + GBP',
             'result': 'IG + FB + GBP'}


def fmt(p):
    if p['kind'] == 'result' and len(p['result']['files']) > 1:
        return 'Carousel'
    return FORMAT[p['kind']]


def tone(p, picks):
    if p['kind'] == 'result':
        return 'Photo'
    photo, _ = ig_photo(p, picks)
    if p['kind'] == 'single' and p['ig'].get('theme') == 'dark':
        return 'Dark'
    return 'Photo' if photo else 'Light'


def credit(pick):
    return f"{pick['by']} / Unsplash" if pick else ''


def write_brief(p, sets, picks):
    folder = os.path.join(ROOT, 'content', p['folder'])
    os.makedirs(os.path.join(folder, 'images'), exist_ok=True)
    keep = os.path.join(folder, 'images', '.gitkeep')
    if not os.path.exists(keep):
        open(keep, 'w').close()
    photo, pick = ig_photo(p, picks)
    gpick = gbp_pick(p, picks)
    g = p['gbp']
    week = (p['_d'] - WEEK_ONE).days // 7 + 1
    files = slide_files(p)
    if p['kind'] == 'result':
        visual = 'Real client result, already designed. Crops to the treatment area, no faces.'
        photo_line = 'Client photos from design/photos/cleo/crops/, shared with written consent.'
    else:
        visual = (f"{len(files)} slides: photo cover, {len(p['ig']['points'])} points, CTA"
                  if p['kind'] == 'carousel' else 'Single image, photo led')
        if p.get('cover_file'):
            visual = 'Carousel: day zero result card as slide 1, then 5 aftercare points and a CTA'
            photo_line = 'Slide 1 is the client day zero card (consent on file). No stock photo.'
        elif p['photo'].get('own'):
            photo_line = ('Using your own photo.' if photo else 'NEEDED: ' + p['photo']['own'])
        elif pick:
            photo_line = (f"{credit(pick)}: https://unsplash.com/photos/{pick['id']} "
                          f"({pick.get('alt') or 'no description'})")
        else:
            photo_line = (f"PENDING. Run design/fetch-octnov-photos.py. It searches Unsplash for "
                          f"\"{p['photo']['ig']}\" and never repeats a photo already used.")
    gbp_photo = (f"{credit(gpick)}: https://unsplash.com/photos/{gpick['id']}" if gpick else
                 f"PENDING. Search: \"{p['photo']['gbp']}\"")
    name = {k: n for k, n, _ in sets}.get(p['set'], '')
    tags = {k: h for k, _, h in sets}.get(p['set'], '')
    text = f"""# {p['title']}

**Post date:** {p['_d']:%A, %B %d, %Y} (Week {week})
**Format:** {fmt(p)}
**Platforms:** {PLATFORMS[p['kind']]}
**Pillar:** {p['pillar']}
**Theme:** {p['theme']}
**SEO keyword (Search Console):** {gsc_label(p['keyword'])}

## Hook
{p['hook']}

## Caption
{p['caption']}

**CTA:** {p['cta']}

## Hashtags ({p['set']}: {name})
{tags}

## Alt text
Paste into Instagram > Advanced settings > Write alt text. Instagram reads it
for search, so it carries the keyword.

{p['alt']}

## Visual / Asset
{visual}

Files in images/: {', '.join(files)}

## Photo
{photo_line}

## Google Business Profile
Card: content/gbp/{p['date']}-{p['slug']}.jpg
Photo: {gbp_photo}

**Short:** {g['text']}

**Long:** in the GBP Daily tab and strategy/gbp-blog-copy.md.

## Design
Style: Editorial (house style)
Template: design/{'carousel' if p['kind'] == 'carousel' else 'single'}.html?post={p['slug']}
Copy lives in strategy/{'plan-' + p['date'][:7]}.json. Edit it there, then run
design/build-octnov.py and design/render.py octnov.

---
Final image(s) for this post are in the images/ subfolder.
"""
    if p['kind'] == 'result':
        text = text.replace(f"Template: design/{'carousel' if p['kind'] == 'carousel' else 'single'}"
                            f".html?post={p['slug']}\n", 'Template: already designed, see design/render.py LATER\n')
    with open(os.path.join(folder, 'brief.md'), 'w', encoding='utf-8') as f:
        f.write(text)


# --- GBP ------------------------------------------------------------------

def update_gbp(posts, picks):
    data = json.load(open(GBP_JSON, encoding='utf-8'))
    photos = json.load(open(GBP_PHOTOS, encoding='utf-8'))
    known = {ph['id'] for ph in photos['photos']}
    for p in posts:
        pk = gbp_pick(p, picks)
        if pk and pk['id'] not in known:
            photos['photos'].append({k: pk[k] for k in ('id', 'by', 'user', 'alt', 'file', 'size')
                                     if k in pk})
            known.add(pk['id'])
    with open(GBP_PHOTOS, 'w', encoding='utf-8') as f:
        json.dump(photos, f, indent=2, ensure_ascii=False)
        f.write('\n')

    kept = [q for q in data['posts'] if q['date'] < START]
    for p in posts:
        g = p['gbp']
        pk = gbp_pick(p, picks)
        kept.append({
            'date': p['date'], 'day': f"{p['_d']:%a}", 'slug': p['slug'],
            'source': f"IG: {p['title']}", 'pillar': p['pillar'],
            'service': g['service'], 'headline': g['headline'], 'script': g['script'],
            'body': g['body'], 'button': g['button'], 'text': g['text'],
            'photo': pk['id'] if pk else '', 'keyword': p['keyword'],
            'description': g['description'],
        })
    data['posts'] = kept
    m = data['_meta']
    m['range'] = f"2026-08-17 to {posts[-1]['date']}"
    m['cadence'] = 'daily through Sep 30; from Oct 1 Monday to Thursday only, matching Instagram'
    m['branch'] = BRANCH
    m['notes'] = [n for n in m['notes'] if not n.startswith('Rows marked source GBP only')]
    for n in ('Rows marked source GBP only are the Wednesdays and Sundays of August and '
              'September, which had no Instagram post. From Oct 1 every GBP post matches that '
              "day's Instagram post, and nothing posts Friday to Sunday.",
              "From Oct 1 each post carries a keyword: the Search Console query it is written "
              "for. It appears in the first sentence of the long description alongside Oakville.",
              'Every photo from Oct 1 is used once. design/fetch-octnov-photos.py refuses any '
              'photo already in strategy/gbp-photos.json or picked for another day.'):
        if n not in m['notes']:
            m['notes'].append(n)
    with open(GBP_JSON, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write('\n')
    print(f'wrote strategy/gbp-daily-content.json  ({len(posts)} posts from {START})')


# --- workbook -------------------------------------------------------------

def update_workbook(posts, picks):
    import openpyxl
    from copy import copy
    from openpyxl.styles import Alignment

    wb = openpyxl.load_workbook(SHEET)
    sets = hashtag_sets(wb)

    ws = wb['Posting Schedule']
    head = [c.value for c in ws[1]]
    for name in SEO_COLS:
        if name not in head:
            c = ws.cell(1, len(head) + 1, name)
            src = ws.cell(1, len(head))
            c.font, c.fill, c.alignment, c.border = (copy(src.font), copy(src.fill),
                                                     copy(src.alignment), copy(src.border))
            head.append(name)
    col = {h: i + 1 for i, h in enumerate(head)}

    first = None
    template = None
    kept, result_meta = {}, {}
    inputs = ['Status', 'Posted Link', 'Reach', 'Saves', 'Comments', 'DMs', 'Bookings']
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not isinstance(d, datetime.datetime):
            continue
        if d.strftime('%Y-%m-%d') < START:
            template = r
            continue
        first = first or r
        kept[d.strftime('%Y-%m-%d')] = {n: ws.cell(r, col[n]).value for n in inputs}
        link = ws.cell(r, col['Consent']).hyperlink
        folder = link.target.rstrip('/').split('/')[-1] if link else None
        if folder:
            result_meta[folder] = (ws.cell(r, col['Photo Assets']).value,
                                   ws.cell(r, col['Consent']).value)
    styles = [ws.cell(template, i) for i in range(1, len(head) + 1)]
    height = ws.row_dimensions[template].height
    if first:
        # these are the last rows in the tab, so nothing below them shifts and
        # no hyperlink can slide out of step with its row
        ws.delete_rows(first, ws.max_row - first + 1)

    for p in posts:
        r = ws.max_row + 1
        photo, pick = ig_photo(p, picks)
        gpick = gbp_pick(p, picks)
        files = slide_files(p)
        folder_url = f"{TREE}/content/{p['folder']}"
        many = len(files) > 1
        img_url = (f"{TREE}/content/{p['folder']}/images" if many else
                   f"{BLOB}/content/{p['folder']}/images/{files[0]}?raw=1")
        gbp_url = f"{BLOB}/{gbp_image(p)}?raw=1"
        stock = p['kind'] != 'result' and not p.get('cover_file')
        # the Instagram post is what the row tracks; its GBP card has its own tab
        done = bool(not stock or photo)
        if p['kind'] == 'result' or p.get('cover_file'):
            old = p['result']['folder']
            assets, consent = result_meta.get(p['folder']) or result_meta.get(old) or ('', 'On file')
        else:
            if p['photo'].get('own'):
                assets = ('Own photo of Cleo' if photo else
                          f"NEEDED: new photo of Cleo, save as design/photos/octnov/{p['date']}-ig.jpg")
            else:
                assets = (f"{credit(pick)} (unsplash.com/photos/{pick['id']})" if pick else
                          f"Photo pending. Unsplash search: {p['photo']['ig']}")
            consent = 'Not needed'
        visual = {'carousel': f"{'READY' if done else 'DESIGNED, photo pending'}: "
                              f"{len(files)} slides in this post's folder",
                  'single': f"{'READY' if done else 'DESIGNED, photo pending'}: "
                            f"final image in this post's folder",
                  'result': f"READY: real client result, "
                            f"{len(files)} {'slides' if many else 'image'} in this post's folder"}[p['kind']]
        k = kept.get(p['date'], {})
        status = k.get('Status')
        if status in (None, '', 'Idea', 'Design', 'Ready'):
            status = 'Ready' if done else 'Design'
        caption = f"{p['caption']}\n\n{dict((a, c) for a, _, c in sets).get(p['set'], '')}"
        values = {
            'Date': datetime.datetime.combine(p['_d'], datetime.time()),
            'Day': f"{p['_d']:%a}", 'Week': (p['_d'] - WEEK_ONE).days // 7 + 1,
            'Month Theme': p['theme'], 'Format': fmt(p), 'Tone': tone(p, picks),
            'Platforms': PLATFORMS[p['kind']], 'Pillar': p['pillar'],
            'Topic & Hook': f"{p['title']} | {p['hook']}", 'Caption Draft': caption,
            'CTA': p['cta'], 'Hashtag Set': p['set'], 'Visual / Asset': visual,
            'Photo Assets': assets, 'Consent': consent, 'Folder Link': 'Open folder',
            'Image Link': f"View {len(files)} slides (jpg)" if many else 'View image (jpg)',
            'Preview': f'=IMAGE("{RAW}/content/{p["folder"]}/images/{files[0]}")',
            'Status': status, 'GBP Caption': p['gbp']['text'],
            'GBP CTA': 'Learn more > Instagram', 'GBP Image Link': 'View GBP image (jpg)',
            'SEO Keyword (GSC)': gsc_label(p['keyword']), 'Alt Text': p['alt'],
        }
        for n in inputs[1:]:
            values[n] = k.get(n)
        for i, name in enumerate(head, start=1):
            c = ws.cell(r, i, values.get(name))
            s = styles[i - 1] if i <= len(styles) else styles[-1]
            c.font, c.fill, c.border = copy(s.font), copy(s.fill), copy(s.border)
            c.alignment, c.number_format = copy(s.alignment), s.number_format
        ws.cell(r, col['Consent']).hyperlink = folder_url
        ws.cell(r, col['Folder Link']).hyperlink = img_url
        ws.cell(r, col['Image Link']).hyperlink = img_url
        ws.cell(r, col['GBP CTA']).hyperlink = gbp_url
        ws.cell(r, col['GBP Image Link']).hyperlink = gbp_url
        for name in ('Caption Draft', 'GBP Caption', 'Alt Text', 'Topic & Hook'):
            ws.cell(r, col[name]).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = height
    ws.column_dimensions[openpyxl.utils.get_column_letter(col['SEO Keyword (GSC)'])].width = 30
    ws.column_dimensions[openpyxl.utils.get_column_letter(col['Alt Text'])].width = 60

    # The status dropdown sat on Folder Link, one column left of Status. Point
    # it at Status, and let it offer the words already typed in the sheet.
    status_col = openpyxl.utils.get_column_letter(col['Status'])
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and 'Idea' in dv.formula1:
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(
                f'{status_col}2:{status_col}{ws.max_row}')
            dv.formula1 = '"Idea,Design,Ready,Scheduled,Published,Hold"'

    move_reels(wb)
    refresh_notes(wb, posts)
    wb.save(SHEET)
    print(f'wrote Posting Schedule  ({len(posts)} rows from {START})')


def hashtag_sets(wb):
    ws = wb['Hashtag Sets']
    rows = [tuple(r[:3]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    if not any(r[0] == 'Set F' for r in rows):
        ws.append(['Set F', 'Electrolysis', SET_F])
        rows.append(('Set F', 'Electrolysis', SET_F))
    return rows


def move_reels(wb):
    """Friday's image or ad slot moves to Thursday, Sunday's reel to Monday."""
    ws = wb['Reels & Video Routine']
    head = [c.value for c in ws[1]]
    di, dy = head.index('Date') + 1, head.index('Day') + 1
    moved = 0
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, di).value
        if not isinstance(d, datetime.datetime) or d.strftime('%Y-%m-%d') < START:
            continue
        shift = {4: -1, 5: 2, 6: 1}.get(d.weekday())
        if shift:
            d = d + datetime.timedelta(days=shift)
            ws.cell(r, di).value = d
            ws.cell(r, dy).value = d.strftime('%a')
            moved += 1
    if moved:
        print(f'moved {moved} Reels & Video slots off Friday and Sunday')


THEMES = [
    ('October', 'Botox & Filler in Oakville',
     'Rank for what people already search: Botox in Oakville, dermal fillers, and '
     'electrolysis, which already sits near the top of Google',
     'First Botox Appointment / Filler Is Not Fake / Botox Treatment Areas / Electrolysis vs Laser / '
     'Botox for Men / Aftercare / Facial Balancing',
     'Real lip results moved onto workdays, electrolysis explainers, luxury studio and Meet Cleo'),
    ('November 1-15', 'Rested by the Holidays',
     'Book now to glow for party season, launch gift certificates, keep electrolysis and '
     'brand searches growing',
     'Holiday Botox Timeline / Party Season Prep / Under Eye Filler',
     'Electrolysis aftercare, gift certificates, Why Cleo Says No, every angle result'),
    ('November 16-30', 'Gift Season',
     'Convert gift and holiday intent into bookings, with honest pricing and electrolysis education',
     'Hormonal Facial Hair / Choosing a Luxury Med Spa / Skincare, Botox or Filler First / '
     'Gift of Glow / December Last Call',
     'Botox areas, lip filler vs lip flip, Black Friday done honestly, does electrolysis hurt'),
]


def refresh_notes(wb, posts):
    ws = wb['Monthly Themes']
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    for period, *vals in THEMES:
        r = rows.get(period) or ws.max_row + 1
        for i, v in enumerate([period, *vals], start=1):
            ws.cell(r, i, v)
            if r > 1 and i <= ws.max_column:
                src = ws.cell(2, i)
                from copy import copy
                ws.cell(r, i).font = copy(src.font)
                ws.cell(r, i).alignment = copy(src.alignment)

    how_to_use(wb['How to Use'], posts)

    ws = wb['Ideas Backlog']
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    used = {'Under-eye filler: the honest truth': 'Scheduled Nov 9',
            'Men and Botox: the fastest growing clients': 'Scheduled Oct 26',
            'Skincare before injectables: the right order': 'Scheduled Nov 23'}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value in used:
            ws.cell(r, 5).value = used[ws.cell(r, 1).value]
    for idea in [
        ('Cheek Filler: The Secret Nobody Talks About', 'Carousel', 'Education',
         'The treatment behind the best natural glow-ups.', 'Dropped from Oct 12 in the SEO replan. Strong for "dermal fillers"'),
        ('Inside the Filler Appointment', 'Single Image', 'Personality & BTS',
         'Planned to the last drop.', 'Dropped from Oct 9 (Friday is off now)'),
        ('5 Filler Lies You Still Believe', 'Carousel', 'Education',
         'The sequel you asked for.', 'Dropped from Oct 26 in the SEO replan'),
        ('3 Months of Honest Answers: The Recap', 'Carousel', 'Trust & Proof',
         'Real education. Real results. Thank you.', 'Dropped from Nov 12. Good for a December recap'),
    ]:
        if idea[0] not in have:
            ws.append(list(idea))


def how_to_use(ws, posts):
    """The prose tab is column B labels and column C text."""
    from copy import copy
    label = {ws.cell(r, 2).value: r for r in range(1, ws.max_row + 1)}
    ws['B3'] = ('Instagram + Google Business Profile content plan. Aug 17 to Nov 30, 2026. '
                'From Oct 1 everything posts Monday to Thursday only.')
    rhythm = label.get('WEEKLY RHYTHM')
    if rhythm:
        rows = [
            ('Monday', 'Carousel, education pillar. The biggest post of the week.'),
            ('Tuesday', 'Single image with a strong hook. Cross-post to Facebook and TikTok.'),
            ('Wednesday', 'From Oct 1: single image or a real client result. Reels slot 1 also lands here.'),
            ('Thursday', 'Carousel: trust, proof or personality pillar, or a real result carousel.'),
            ('Fri to Sun', 'Off. From Oct 1 nothing is scheduled Friday, Saturday or Sunday, '
                           'on any tab.'),
        ]
        for i, (day, text) in enumerate(rows, start=rhythm + 1):
            ws.cell(i, 2).value, ws.cell(i, 3).value = day, text
    reels = label.get('REELS & VIDEO ROUTINE')
    if reels:
        for r in range(reels + 1, reels + 9):
            b = ws.cell(r, 2).value
            if b == 'Friday':
                ws.cell(r, 3).value = ('Slot 2. Image post or paid ad creative. From Oct 1 it '
                                       'moves to the Thursday of the same week.')
            elif b == 'Wednesday':
                ws.cell(r, 3).value = ('Slot 1. Reel. From Oct 1 Wednesday also has a feed post, '
                                       'so the reel goes up alongside it.')
            elif b == 'Sunday':
                ws.cell(r, 3).value = 'Slot 3. Reel. From Oct 1 it moves to the Monday after.'
            elif b == 'New themes':
                ws.cell(r, 3).value = ('Posting Schedule now runs to Nov 30 (Oct: Botox & Filler '
                                       'in Oakville, Nov 1-15: Rested by the Holidays, Nov 16-30: '
                                       'Gift Season). The routine carries on to January.')
    if 'SEARCH CONSOLE (SEO)' in label:
        return
    head, lab, txt = ws['B11'], ws['B12'], ws['C12']
    r = ws.max_row + 2

    def put(row, b, c=None, style=(lab, txt)):
        ws.cell(row, 2, b)
        ws.cell(row, 2).font = copy(style[0].font)
        ws.cell(row, 2).alignment = copy(style[0].alignment)
        if c is not None:
            ws.cell(row, 3, c)
            ws.cell(row, 3).font = copy(style[1].font)
            ws.cell(row, 3).alignment = copy(txt.alignment)

    put(r, 'SEARCH CONSOLE (SEO)', style=(head, head))
    put(r + 1, 'Source', 'Search Console export, last 28 days to Sep 28, 2026. Every post from '
                         'Oct 1 is written for one query, named in the SEO Keyword (GSC) column.')
    by_kw = {}
    for p in posts:
        by_kw.setdefault(p['keyword'], []).append(p['_d'].strftime('%b %d'))
    row = r + 2
    for kw, dates in by_kw.items():
        put(row, gsc_label(kw), f"{len(dates)} posts: {', '.join(dates)}")
        row += 1
    put(row, 'Alt text', 'Paste the Alt Text column into Instagram > Advanced settings > '
                         'Write alt text before posting. Instagram search reads it.')


def main():
    posts = load_plan()
    picks = load_picks()
    move_results(posts)
    prune_old(posts)
    prune_gbp(posts)
    write_data_js(posts, picks)
    import openpyxl
    sets = hashtag_sets(openpyxl.load_workbook(SHEET))
    for p in posts:
        write_brief(p, sets, picks)
    print(f'wrote {len(posts)} briefs')
    update_gbp(posts, picks)
    update_workbook(posts, picks)
    for script in ('design/build-gbp-daily.py', 'strategy/organize-workbook.py',
                   'design/build-blog-copy.py'):
        subprocess.run([sys.executable, os.path.join(ROOT, script)], check=True)
    pending = sum(1 for p in posts if not gbp_pick(p, picks)) + sum(
        1 for p in posts if p['kind'] in ('carousel', 'single') and not p.get('cover_file')
        and p['photo'].get('ig') and not ig_photo(p, picks)[0])
    own = [p['date'] for p in posts if p['photo'].get('own') and not ig_photo(p, picks)[0]]
    if own:
        print(f"\nwaiting on your own photo of Cleo for: {', '.join(own)}")
    if pending:
        print(f'\n{pending} photos still pending. Run design/fetch-octnov-photos.py, '
              'then this script, then design/render.py octnov.')
    print('\nnext: python3 design/render.py octnov')


if __name__ == '__main__':
    main()
